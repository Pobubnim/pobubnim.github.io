# -*- coding: utf-8 -*-
"""Приёмка вызывного листа v2 (instrumenty/vyzyvnoj-list.html) на живой странице.

Проверяет всё, чем лист отличается от бумажного шаблона: свет по координатам
(сверка с astral), свободное расписание с заготовками и сдвигом смены, ссылки
на карту, память людей и вставку списком, блоки по выбору, сцены из шот-листа,
готовность листа, рассылку (общий и личный текст, wa.me, статусы, «что
изменилось»), календарь .ics, таблицу .csv, .docx с распаковкой, печать в PDF,
черновик с переводом листов первой версии и мобильную раскладку.

Запуск:  python tools/test_callsheet.py [url]
"""
import base64
import datetime
import io
import json
import re
import subprocess
import sys
import tempfile
import time
import urllib.parse
import urllib.request
import zipfile
import zoneinfo
from xml.dom import minidom

import websocket

try:
    import cv2
    import numpy as np
except ImportError:
    cv2 = None

try:
    from astral import LocationInfo
    from astral.sun import sun as astral_sun
except ImportError:
    astral_sun = None

CHROME = r"C:\Program Files\Google\Chrome\Application\chrome.exe"
PORT = 9391
URL = sys.argv[1] if len(sys.argv) > 1 else "http://localhost:8765/instrumenty/vyzyvnoj-list.html"
KEY = "pobubnim-callsheet-v1"

fails = []


def check(name, cond, got=""):
    print(("ok  " if cond else "ФЕЙЛ ") + name + ("" if cond else "  -> " + str(got)[:300]))
    if not cond:
        fails.append(name)


class Tab:
    def __init__(self, width=1280):
        self.proc = subprocess.Popen(
            [CHROME, "--headless=new", "--disable-gpu", "--hide-scrollbars",
             "--user-data-dir=" + tempfile.mkdtemp(prefix="pbchrome-"),
             f"--remote-debugging-port={PORT}", "--remote-allow-origins=*",
             f"--window-size={width},900", "about:blank"],
            stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        tabs = None
        for _ in range(80):
            try:
                allt = json.load(urllib.request.urlopen(f"http://127.0.0.1:{PORT}/json"))
                tabs = [t for t in allt if t.get("type") == "page"]
                if tabs:
                    break
            except Exception:
                pass
            time.sleep(0.25)
        if not tabs:
            raise RuntimeError("Chrome не поднялся")
        self.ws = websocket.create_connection(tabs[0]["webSocketDebuggerUrl"], timeout=60)
        self.mid = 0
        self.errors = []
        self.cmd("Runtime.enable")
        self.cmd("Page.enable")
        # худший случай раскладки: системный шрифт вместо Inter
        self.cmd("Network.enable")
        self.cmd("Network.setBlockedURLs", urls=["*fonts.googleapis.com*", "*fonts.gstatic.com*"])

    def cmd(self, method, **params):
        self.mid += 1
        self.ws.send(json.dumps({"id": self.mid, "method": method, "params": params}))
        while True:
            msg = json.loads(self.ws.recv())
            if msg.get("id") == self.mid:
                return msg
            if msg.get("method") == "Runtime.exceptionThrown":
                d = msg["params"]["exceptionDetails"]
                self.errors.append((d.get("exception") or {}).get("description") or d.get("text", "JS error"))

    def js(self, expr, wait=False):
        r = self.cmd("Runtime.evaluate", expression=expr, returnByValue=True, awaitPromise=wait)
        res = r.get("result", {})
        if "exceptionDetails" in res:
            d = res["exceptionDetails"]
            self.errors.append((d.get("exception") or {}).get("description") or str(d.get("text")))
            return None
        return res.get("result", {}).get("value")

    def set(self, el_id, value):
        self.js("(()=>{const e=document.getElementById('%s');e.value=%s;"
                "e.dispatchEvent(new Event('input',{bubbles:true}));"
                "e.dispatchEvent(new Event('change',{bubbles:true}));})()" % (el_id, json.dumps(value)))

    def row(self, lst, i, k, value):
        """Вписать значение в строку списка так, как это делает человек: input, потом change."""
        self.js("(()=>{const e=document.querySelectorAll('#%s .row-i')[%d].querySelector('[data-k=\"%s\"]');"
                "e.value=%s;e.dispatchEvent(new Event('input',{bubbles:true}));"
                "e.dispatchEvent(new Event('change',{bubbles:true}));})()" % (lst, i, k, json.dumps(value)))

    def click(self, sel):
        self.js("document.querySelector(%s).click()" % json.dumps(sel))

    def paper(self):
        return self.js("document.getElementById('paper').innerText") or ""

    def goto(self, url):
        self.cmd("Page.navigate", url=url)
        for _ in range(30):
            time.sleep(0.4)
            if self.js("typeof window.PobubnimCallsheet === 'object' && typeof PobubnimCallsheet.chatText") == "function":
                return
        raise RuntimeError("страница не поднялась: " + "; ".join(self.errors[:2]))

    def close(self):
        try:
            self.ws.close()
        finally:
            self.proc.kill()


def diff_min(a, b):
    ah, am = (int(x) for x in a.split(":"))
    bh, bm = (int(x) for x in b.split(":"))
    return abs((ah * 60 + am) - (bh * 60 + bm))


def sched(t):
    return json.loads(t.js("JSON.stringify(PobubnimCallsheet.state().sched.map(r=>[r.t,r.what]))"))


def main():
    t = Tab()
    try:
        t.goto(URL)

        # ---------- 1. пустой лист: подсказка первого визита и готовность ----------
        check("пустой лист показывает подсказку первого визита", t.js("!document.getElementById('starter').hidden"))
        check("готовность считает с нуля",
              t.js("document.getElementById('ready-n').textContent").startswith("Готово 0 из"),
              t.js("document.getElementById('ready-n').textContent"))
        check("пустой лист не кладётся в хранилище", t.js(f"localStorage.getItem('{KEY}')") in (None, ""))

        # ---------- 2. город, координаты и свет ----------
        t.set("f-city", "Наро-Фоминск")
        lat = t.js("document.getElementById('f-lat').value")
        tz = t.js("document.getElementById('f-tz').value")
        check("город подставил координаты", lat.startswith("55.38") and tz == "3", f"{lat} / {tz}")
        t.set("f-date", "2026-09-14")
        paper = t.paper()
        got = json.loads(t.js("JSON.stringify((function(){var s=PobubnimCallsheet.sun();"
                              "return [PobubnimSun.hhmm(s.sunrise),PobubnimSun.hhmm(s.sunset)];})())"))
        check("в листе появился блок света", "Восход и закат" in paper and "Золотой час" in paper, paper[:80])
        if astral_sun:
            ref = astral_sun(LocationInfo("НФ", "RU", "Europe/Moscow", 55.3853, 36.7325).observer,
                             date=datetime.date(2026, 9, 14), tzinfo=zoneinfo.ZoneInfo("Europe/Moscow"))
            rise, sset = ref["sunrise"].strftime("%H:%M"), ref["sunset"].strftime("%H:%M")
            check("восход совпал с эталоном astral", diff_min(got[0], rise) <= 2, f"{got[0]} против {rise}")
            check("закат совпал с эталоном astral", diff_min(got[1], sset) <= 2, f"{got[1]} против {sset}")
        check("закат стоит в плитке листа", got[1] in t.js("document.querySelector('#paper table.cs-facts').innerText"), got[1])
        check("дата на листе с днём недели", "понедельник, 14 сентября 2026" in paper, paper[:160])
        t.set("f-lat", "55.7558, 37.6173")
        check("точка из карт разложилась на широту и долготу",
              t.js("document.getElementById('f-lat').value") == "55.7558" and
              t.js("document.getElementById('f-lng').value") == "37.6173",
              t.js("document.getElementById('f-lat').value + ' / ' + document.getElementById('f-lng').value"))
        t.set("f-city", "Наро-Фоминск")

        # ---------- 3. расписание ----------
        t.row("sched", 0, "t", "07:00")
        t.row("sched", 1, "t", "09:00")
        t.row("sched", 2, "t", "22:00")
        paper = t.paper()
        facts = t.js("document.querySelector('#paper table.cs-facts').innerText")
        check("плитки: сбор, начало, конец", all(x in facts for x in ("07:00", "09:00", "22:00")), facts)
        check("расписание в листе с длительностью", "Сбор группы" in paper and "2 ч" in paper and "Смена целиком" in paper, "")
        check("предупреждение про конец после заката", "после заката" in paper, "")
        check("подсказка про золотой час в смене", "попадает в смену" in paper, "")
        check("полоса дня: день, золотой и синий час, смена, точки",
              t.js("['day','gold','blue','shift','tick'].every(c=>document.querySelector('#sunbar .track .'+c))"),
              t.js("document.querySelector('#sunbar .track').innerHTML")[:200])
        check("длинная смена даёт подсказку про переработку",
              "переработке" in t.js("document.getElementById('ready-warn').textContent"),
              t.js("document.getElementById('ready-warn').textContent"))

        # заготовка от времени сбора
        t.js("[...document.querySelectorAll('#presets .chip')].find(b=>b.textContent==='Интервью').click()")
        s = sched(t)
        check("заготовка «Интервью» встала от времени сбора",
              len(s) == 6 and s[0] == ["07:00", "Сбор группы"] and s[3][0] == "08:30" and s[5] == ["11:00", "Конец смены"], s)
        t.click(".pb-toast.on button")
        check("заготовка отменяется кнопкой «Вернуть»", len(sched(t)) == 3 and sched(t)[2][0] == "22:00", sched(t))

        # сдвиг смены: всё время в листе уезжает разом
        t.row("locs", 0, "name", "Усадьба")
        t.row("locs", 0, "addr", "ул. Ленина, 14")
        t.row("locs", 0, "time", "09:00")
        t.row("crew", 0, "who", "Иван Петров")
        t.row("crew", 0, "at", "06:30")
        t.click('#shift [data-d="15"]')
        st = json.loads(t.js("JSON.stringify((function(){var S=PobubnimCallsheet.state();"
                             "return [S.sched[0].t,S.sched[2].t,S.locs[0].time,S.crew[0].at];})())"))
        check("сдвиг +15 двигает расписание, локации и личные вызовы", st == ["07:15", "22:15", "09:15", "06:45"], st)
        t.js("document.getElementById('shift-from').value='2'")
        t.click('#shift [data-d="-15"]')
        st = json.loads(t.js("JSON.stringify((function(){var S=PobubnimCallsheet.state();"
                             "return [S.sched[0].t,S.sched[2].t,S.crew[0].at];})())"))
        check("сдвиг с выбранной строки не трогает то, что раньше неё", st == ["07:15", "22:00", "06:45"], st)
        t.js("document.getElementById('shift-from').value='0'")
        for i, v in enumerate(("07:00", "09:00", "22:00")):
            t.row("sched", i, "t", v)
        t.row("crew", 0, "at", "06:30")
        t.row("locs", 0, "time", "09:00")

        # порядок времени
        t.click("#add-sched")
        t.row("sched", 3, "t", "13:00")
        t.row("sched", 3, "what", "Обед")
        check("строки не по порядку замечены", "не по порядку" in t.js("document.getElementById('ready-warn').textContent"), "")
        t.click("#btn-sort")
        s = sched(t)
        check("«Расставить по времени» упорядочила строки", [r[0] for r in s] == ["07:00", "09:00", "13:00", "22:00"], s)
        t.row("sched", 3, "t", "02:00")
        k = json.loads(t.js("JSON.stringify((function(){var k=PobubnimCallsheet.key();return [k.tl.messy,k.end&&k.end.m-k.call.m];})())"))
        check("смена через полночь считается, а не ругается", k == [False, 19 * 60], k)
        t.row("sched", 3, "t", "22:00")
        check("Enter на последней строке заводит новую",
              t.js("(()=>{const e=document.querySelectorAll('#sched .row-i')[3].querySelector('[data-k=\"what\"]');"
                   "e.dispatchEvent(new KeyboardEvent('keydown',{key:'Enter',bubbles:true}));"
                   "return document.querySelectorAll('#sched .row-i').length})()") == 5, "")
        t.js("document.querySelectorAll('#sched .row-i')[4].querySelector('.del-row').click()")

        # ---------- 4. локации ----------
        href = t.js("document.querySelector('#paper a[href*=\"yandex.ru/maps\"]').href")
        check("адрес ведёт на карту, город дописан",
              urllib.parse.unquote(href) == "https://yandex.ru/maps/?text=Наро-Фоминск, ул. Ленина, 14", urllib.parse.unquote(href or ""))
        t.click("#add-loc")
        t.row("locs", 1, "name", "Студия")
        check("две локации в листе", "Усадьба" in t.paper() and "Студия" in t.paper(), "")
        t.js("document.querySelectorAll('#locs .row-i')[0].querySelector('.down').click()")
        check("локации переставляются", t.js("PobubnimCallsheet.state().locs[0].name") == "Студия", "")
        t.js("document.querySelectorAll('#locs .row-i')[0].querySelector('.del-row').click()")
        check("локация удаляется", t.js("PobubnimCallsheet.state().locs.length") == 1, "")
        t.click(".pb-toast.on button")
        check("удалённая строка возвращается на своё место",
              t.js("PobubnimCallsheet.state().locs.map(l=>l.name).join('|')") == "Студия|Усадьба",
              t.js("PobubnimCallsheet.state().locs.map(l=>l.name).join('|')"))
        t.js("document.querySelectorAll('#locs .row-i')[0].querySelector('.del-row').click()")

        # ---------- 5. группа ----------
        t.js("[...document.querySelectorAll('#chips .chip')].find(b=>b.textContent==='Звукорежиссёр').click()")
        check("роль добавлена чипом", t.js("PobubnimCallsheet.state().crew.map(p=>p.role).join('|')") ==
              "Оператор-постановщик|Звукорежиссёр", t.js("PobubnimCallsheet.state().crew.map(p=>p.role).join('|')"))
        t.row("crew", 0, "phone", "89161234567")
        check("телефон приведён к виду +7 ···", t.js("PobubnimCallsheet.state().crew[0].phone") == "+7 916 123-45-67",
              t.js("PobubnimCallsheet.state().crew[0].phone"))
        check("телефон на листе набирается по клику",
              t.js("document.querySelector('#paper a[href=\"tel:+79161234567\"]').textContent") == "+7 916 123-45-67", "")
        t.row("crew", 1, "who", "Олег Жуков")
        t.row("crew", 1, "phone", "+7 925 000-11-22")
        paper = t.paper()
        check("человек без своего времени вызван к общему сбору", re.search(r"Олег Жуков\s+\+7 925 000-11-22\s+07:00", paper), paper[-400:])
        check("на листе есть «На связи»", "На связи в день съёмки: Иван Петров" in paper, "")

        parsed = json.loads(t.js("JSON.stringify(PobubnimCallsheet.parsePeople("
                                 + json.dumps("1. Гафер — Игорь Лапин — 8 916 000-00-01 — к 07:00\n"
                                              "Мария Соколова, продюсер, +7 925 111-22-33\n"
                                              "Аня\n\n- визажист: Даша 9:15") + "))"))
        want = [["Гафер", "Игорь Лапин", "+7 916 000-00-01", "07:00"], ["Продюсер", "Мария Соколова", "+7 925 111-22-33", ""],
                ["", "Аня", "", ""], ["Визажист", "Даша", "", "09:15"]]
        check("вставка списком разбирает четыре формата строки",
              [[p["role"], p["who"], p["phone"], p["at"]] for p in parsed] == want,
              [[p["role"], p["who"], p["phone"], p["at"]] for p in parsed])
        t.js("document.getElementById('paste').open=true")
        t.set("paste-text", "Гафер — Игорь Лапин — 8 916 000-00-01 — к 07:00\nМария Соколова, продюсер, +7 925 111-22-33")
        t.click("#paste-go")
        check("список добавил людей в группу", t.js("PobubnimCallsheet.state().crew.length") == 4, "")
        check("выбор «кто на связи» появился и меняет лист",
              t.js("!document.getElementById('contact-box').hidden"), "")
        t.set("f-contact", "Мария Соколова")
        check("на связи — выбранный человек", "На связи в день съёмки: Мария Соколова (продюсер)" in t.paper(), "")

        # память людей: знакомое имя подставляет телефон и роль
        time.sleep(0.7)
        t.click("#add-crew")
        t.row("crew", 4, "who", "Игорь Лапин")
        mem = json.loads(t.js("JSON.stringify((function(){var p=PobubnimCallsheet.state().crew[4];return [p.phone,p.role];})())"))
        check("знакомое имя подставило телефон и роль", mem == ["+7 916 000-00-01", "Гафер"], mem)
        t.js("document.querySelectorAll('#crew .row-i')[4].querySelector('.del-row').click()")

        # большая группа — по цехам
        for name in ("Анна", "Борис", "Вера"):
            t.click("#add-crew")
            n = t.js("PobubnimCallsheet.state().crew.length") - 1
            t.row("crew", n, "who", name)
            t.row("crew", n, "role", "Осветитель")
        check("группа от семи человек собрана по цехам",
              t.js("[...document.querySelectorAll('#paper tr.grp')].map(r=>r.textContent).join('|')") ==
              "Продакшн|Камера|Свет|Звук", t.js("[...document.querySelectorAll('#paper tr.grp')].map(r=>r.textContent).join('|')"))

        # ---------- 6. блоки по выбору ----------
        t.click('#block-chips [data-b="cast"]')
        t.row("cast", 0, "who", "Тимур Асланов")
        t.row("cast", 0, "part", "бариста")
        t.row("cast", 0, "at", "08:30")
        t.row("cast", 0, "ready", "09:00")
        t.click('#block-chips [data-b="safety"]')
        t.set("f-hosp", "Травмпункт")
        t.set("f-hospaddr", "ул. Больничная, 1")
        t.click('#block-chips [data-b="food"]')
        t.set("t-food", "Обед в 13:00 на площадке.")
        paper = t.paper()
        check("блок «В кадре» на листе", "В КАДРЕ" in paper.upper() and "бариста" in paper, "")
        check("безопасность: 112 и ближайшая помощь", "112" in paper and "Травмпункт" in paper, "")
        check("текстовый блок «Питание» на листе", "Обед в 13:00 на площадке." in paper, "")
        t.click('#block-chips [data-b="food"]')
        check("выключенный блок уходит с листа, текст в нём остаётся",
              "Обед в 13:00 на площадке." not in t.paper() and t.js("PobubnimCallsheet.state().txt.food") == "Обед в 13:00 на площадке.", "")

        # сцены из шот-листа
        t.js("localStorage.setItem('pobubnim-shotlist-v1', JSON.stringify({defmin:'10',scenes:["
             "{name:'Утро в цехе',loc:'Цех',time:'день',shots:[{what:'Общий план',min:'15'},{what:'Деталь'}]},"
             "{name:'',loc:'',time:'день',shots:[{what:''}]}]}))")
        t.click('#block-chips [data-b="scenes"]')
        check("кнопка «из шот-листа» появилась", t.js("!document.getElementById('btn-shots').hidden"), "")
        t.click("#btn-shots")
        sc = json.loads(t.js("JSON.stringify(PobubnimCallsheet.state().scenes)"))
        check("сцены подтянулись из шот-листа с объёмом",
              len(sc) == 1 and sc[0]["what"] == "Утро в цехе" and sc[0]["where"] == "Цех, день" and sc[0]["note"] == "2 кадра · 25 мин", sc)

        # ---------- 7. версия для клиента ----------
        t.js("(()=>{const c=document.getElementById('f-nophones');c.checked=true;c.dispatchEvent(new Event('change',{bubbles:true}));})()")
        paper = t.paper()
        check("версия для клиента: телефонов группы нет, контакт остался",
              "+7 925 000-11-22" not in paper and "+7 925 111-22-33" in paper and "Телефон" not in paper, "")
        check("версия для клиента: телефонов нет и в тексте для чата",
              "+7 925 000-11-22" not in t.js("PobubnimCallsheet.chatText()"), "")
        t.js("(()=>{const c=document.getElementById('f-nophones');c.checked=false;c.dispatchEvent(new Event('change',{bubbles:true}));})()")

        # ---------- 8. готовность ----------
        ready = t.js("document.getElementById('ready-n').textContent + ' / ' + document.getElementById('ready-miss').textContent")
        check("готовность называет, чего не хватает", "Готово 6 из 8" in ready and "название проекта" in ready and "телефон у 3 человек" in ready, ready)
        t.js("[...document.querySelectorAll('#ready-miss .miss')].find(b=>b.textContent==='название проекта').click()")
        check("клик по недостающему ставит курсор в поле", t.js("document.activeElement.id") == "f-proj", t.js("document.activeElement.id"))
        t.set("f-proj", "Рекламный ролик EVERGO")
        t.set("f-day", "1 из 2")

        # ---------- 9. рассылка ----------
        chat = t.js("PobubnimCallsheet.chatText()")
        check("общий текст: шапка, времена, карта, группа",
              chat.startswith("ВЫЗЫВНОЙ ЛИСТ · «Рекламный ролик EVERGO», смена 1 из 2") and "Общий сбор — 07:00" in chat and
              "Карта: https://yandex.ru/maps/?text=" in chat and "Гафер — Игорь Лапин, +7 916 000-00-01 — к 07:00" in chat and
              "На связи: Мария Соколова (продюсер)" in chat, chat[:400])
        SIGN = "Лист собран в бесплатном конструкторе: https://pobubnim.ru/instrumenty/vyzyvnoj-list.html"
        check("общий текст кончается строкой со ссылкой на конструктор", chat.rstrip().endswith(SIGN), chat[-160:])
        t.click("#btn-send")
        check("галочка подписи в окне рассылки стоит, строка видна в тексте",
              t.js("document.getElementById('f-sign').checked") is True and SIGN in t.js("document.getElementById('msg-all').textContent"), "")
        t.js("(()=>{const c=document.getElementById('f-sign');c.checked=false;c.dispatchEvent(new Event('change',{bubbles:true}));})()")
        time.sleep(0.4)
        check("без галочки подписи нет ни в тексте окна, ни в ссылке мессенджера",
              SIGN not in t.js("document.getElementById('msg-all').textContent") and
              "pobubnim.ru%2Finstrumenty" not in t.js("document.getElementById('btn-wa').href") and
              t.js("PobubnimCallsheet.state().noSign") is True, t.js("document.getElementById('msg-all').textContent")[-120:])
        t.js("(()=>{const c=document.getElementById('f-sign');c.checked=true;c.dispatchEvent(new Event('change',{bubbles:true}));})()")
        time.sleep(0.4)
        check("галочка вернула подпись", SIGN in t.js("document.getElementById('msg-all').textContent") and
              "pobubnim.ru%2Finstrumenty" in t.js("document.getElementById('btn-wa').href"), "")
        t.js("document.getElementById('send').close()")
        check("подпись листа — ссылка на сам инструмент",
              t.js("(document.querySelector('#paper .doc-note a.pb-by')||{}).href") == "https://pobubnim.ru/instrumenty/vyzyvnoj-list.html", "")
        per = t.js("PobubnimCallsheet.personal(PobubnimCallsheet.state().crew[0],'crew')")
        check("личный вызов: имя, своё время, адрес и карта",
              per.startswith("Иван, вызывной лист на понедельник, 14 сентября.") and "Вызов: 06:30 · Оператор-постановщик" in per and
              "Куда: Усадьба, ул. Ленина, 14" in per and "Карта: https://" in per and "Напишите, пожалуйста, что получили." in per, per)
        t.click("#btn-send")
        check("окно рассылки открылось", t.js("document.getElementById('send').open"), "")
        check("в окне по строке на человека", t.js("document.querySelectorAll('#per .per-i').length") == 8,
              t.js("document.querySelectorAll('#per .per-i').length"))
        wa = t.js("document.querySelector('#per .per-i a[data-act=\"wa\"]').href")
        check("ссылка WhatsApp: номер цифрами и личный текст",
              wa.startswith("https://wa.me/79161234567?text=") and "Иван" in urllib.parse.unquote(wa), wa[:80])
        tg = urllib.parse.unquote(t.js("document.getElementById('btn-tg').href"))
        check("ссылка Telegram: url обязателен и ведёт на карту", tg.startswith("https://t.me/share/url?url=https://yandex.ru/maps/?text="), tg[:90])
        check("длинный лист честно говорит про краткий вариант в ссылке",
              t.js("document.getElementById('url-note').hidden") == (len(urllib.parse.quote(chat, safe="")) <= 6000), "")
        t.js("(()=>{const s=document.querySelectorAll('#per .per-st')[0];s.value='ok';s.dispatchEvent(new Event('change',{bubbles:true}));})()")
        t.js("(()=>{const s=document.querySelectorAll('#per .per-st')[1];s.value='sent';s.dispatchEvent(new Event('change',{bubbles:true}));})()")
        check("статусы считаются", t.js("document.getElementById('per-sum').textContent") == "Отправлено 2 из 8 · подтвердили 1",
              t.js("document.getElementById('per-sum').textContent"))
        check("напоминание доступно, пока не все подтвердили", t.js("!document.getElementById('btn-remind').hidden"), "")

        # версия листа и «что изменилось»
        t.click("#btn-mark")
        check("отметка «разослан» запомнила версию 1", t.js("PobubnimCallsheet.state().rev") == 1 and
              t.js("document.getElementById('btn-mark').disabled") and t.js("document.getElementById('diff-box').hidden"), "")
        t.row("sched", 0, "t", "07:30")
        t.row("locs", 0, "addr", "ул. Мира, 5")
        t.row("crew", 1, "who", "Олег Жуковский")
        ch = json.loads(t.js("JSON.stringify(PobubnimCallsheet.changes())"))
        check("изменения после рассылки названы по именам",
              "Сбор группы: 07:00 → 07:30" in ch and "Локация 1: теперь Усадьба, ул. Мира, 5" in ch and
              "Не участвует: Олег Жуков" in ch and any(x.startswith("В группе: Олег Жуковский") for x in ch), ch)
        check("окно показывает изменения списком", t.js("!document.getElementById('diff-box').hidden") and
              t.js("document.querySelectorAll('#diff li').length") == len(ch), t.js("document.getElementById('diff').innerHTML"))
        txt = t.js("PobubnimCallsheet.changesText()")
        check("сообщение об изменениях: шапка, пункты, номер версии",
              txt.startswith("ИЗМЕНЕНИЯ В ВЫЗЫВНОМ · «Рекламный ролик EVERGO», смена 1 из 2 · 14 сентября") and
              "— Сбор группы: 07:00 → 07:30" in txt and txt.endswith("Актуальная версия: 2."), txt)
        t.click("#btn-mark")
        check("вторая рассылка ставит на лист «ВЕРСИЯ 2»", "ВЕРСИЯ 2" in t.paper(), t.paper()[:60])
        t.click("#send-x")
        check("окно закрывается", not t.js("document.getElementById('send').open"), "")
        t.row("sched", 0, "t", "07:00")

        # ---------- 9б. код маршрута, логотип, прогноз ----------
        src = t.js("(document.querySelector('#paper .cs-qr img')||{}).src") or ""
        want_url = t.js("PobubnimCallsheet.mapUrl(PobubnimCallsheet.state().locs[0].addr)")
        check("у локации с адресом на листе стоит QR-код", src.startswith("data:image/png;base64,"), src[:40])
        if cv2 is not None and src:
            img = cv2.imdecode(np.frombuffer(base64.b64decode(src.split(",", 1)[1]), np.uint8), cv2.IMREAD_GRAYSCALE)
            got_url = cv2.QRCodeDetector().detectAndDecode(cv2.resize(img, None, fx=2, fy=2, interpolation=cv2.INTER_NEAREST))[0]
            check("QR читается независимым декодером и ведёт на карту", got_url == want_url, f"{got_url!r} против {want_url!r}")
        long_ok = t.js("(()=>{const q=PobubnimQR.matrix('x'.repeat(858));return q&&q.version===20&&PobubnimQR.matrix('x'.repeat(859))===null})()")
        check("кодировщик берёт ровно 858 байт и честно отказывает на 859", long_ok, "")
        t.click('#block-chips [data-b="qr"]')
        check("чип «QR маршрута» убирает код с листа", t.js("!document.querySelector('#paper .cs-qr')"), "")
        t.click('#block-chips [data-b="qr"]')

        logo_file = tempfile.mktemp(suffix=".png")
        if cv2 is not None:
            pic = np.full((300, 900, 3), 255, np.uint8)
            cv2.rectangle(pic, (40, 40), (860, 260), (30, 30, 30), -1)
            cv2.imwrite(logo_file, pic)
            node = t.cmd("DOM.getDocument")["result"]["root"]["nodeId"]
            inp = t.cmd("DOM.querySelector", nodeId=node, selector="#f-logo")["result"]["nodeId"]
            t.cmd("DOM.setFileInputFiles", nodeId=inp, files=[logo_file])
            time.sleep(0.8)
            lg = json.loads(t.js("JSON.stringify((function(){var i=document.querySelector('#paper .cs-logo img');"
                                 "return i?[i.getAttribute('width'),i.getAttribute('height'),i.src.slice(0,22)]:null})())"))
            check("логотип встал на лист, приведён к 320 точкам по ширине", lg == ["160", "54", "data:image/png;base64,"], lg)
            check("кнопка логотипа сменила подпись, появилось «Убрать»",
                  t.js("document.getElementById('btn-logo').textContent") == "Заменить логотип" and
                  t.js("!document.getElementById('btn-logo-del').hidden"), "")

        t.js("""(()=>{const rows=[];for(let h=0;h<24;h++){const wet=h===14||h===15;
          rows.push({time:'2026-09-14T'+String(h).padStart(2,'0')+':00:00Z',data:{instant:{details:{
            air_temperature:h<4||h>19?-5:10+(h-4)*0.5,wind_speed:h<4||h>19?20:(h===12?5.6:3.4)}},
            next_1_hours:{summary:{symbol_code:wet?'lightrain':'partlycloudy_day'},details:{precipitation_amount:wet?0.6:0}}}});}
          window.__wx={n:0,url:''};
          const real=window.fetch;
          window.fetch=(u,o)=>{if(String(u).indexOf('api.met.no')<0)return real(u,o);   /* счётчик Метрики тоже зовёт fetch */
            window.__wx.n++;window.__wx.url=u;return Promise.resolve({ok:true,
            headers:{get:()=>new Date(Date.now()+3600e3).toUTCString()},json:()=>Promise.resolve({properties:{timeseries:rows}})});};
          sessionStorage.clear();})()""")
        t.click("#btn-wx")
        time.sleep(0.5)
        wx = t.js("document.getElementById('f-wx').value")
        check("прогноз посчитан за часы смены, а не за сутки",
              wx == "+10…+18, переменная облачность, небольшой дождь, около 1 мм, ветер до 6 м/с — прогноз MET Norway", wx)
        check("в сервис ушли только координаты, не больше четырёх знаков",
              t.js("window.__wx.url") == "https://api.met.no/weatherapi/locationforecast/2.0/compact?lat=55.3853&lon=36.7325",
              t.js("window.__wx.url"))
        check("прогноз попал на лист", "Погода" in t.paper() and "прогноз MET Norway" in t.paper(), "")
        t.click("#btn-wx")
        time.sleep(0.4)
        check("повторный запрос до срока Expires не уходит в сеть", t.js("window.__wx.n") == 1, t.js("window.__wx.n"))
        t.set("f-date", "2026-11-30")
        t.click("#btn-wx")
        time.sleep(0.4)
        check("дата за горизонтом прогноза названа прямо",
              "на эту дату его пока нет" in (t.js("document.querySelector('.pb-toast').textContent") or ""),
              t.js("document.querySelector('.pb-toast').textContent"))
        t.js("sessionStorage.clear();window.fetch=(u)=>String(u).indexOf('api.met.no')<0?Promise.resolve({ok:true}):Promise.reject(new Error('сеть'))")
        t.click("#btn-wx")
        time.sleep(0.4)
        check("сбой сервиса не ломает лист: поле осталось, человеку сказано, что делать",
              "не ответил" in (t.js("document.querySelector('.pb-toast').textContent") or "") and
              t.js("document.getElementById('f-wx').value") == wx and not t.js("document.getElementById('btn-wx').disabled"), "")
        t.set("f-date", "2026-09-14")

        # ---------- 10. календарь и таблица ----------
        ics = t.js("PobubnimCallsheet.ics()")
        lines = ics.split("\r\n")
        check(".ics: каркас и время в UTC (07:00 по Москве = 04:00Z)",
              lines[0] == "BEGIN:VCALENDAR" and lines[-2] == "END:VCALENDAR" and "DTSTART:20260914T040000Z" in lines and
              "DTEND:20260914T190000Z" in lines and "BEGIN:VALARM" in lines, lines[:12])
        check(".ics: строки не длиннее 75 октетов", max(len(x.encode("utf-8")) for x in lines) <= 75,
              max(len(x.encode("utf-8")) for x in lines))
        unfolded = ics.replace("\r\n ", "")
        check(".ics: описание разворачивается обратно без потерь",
              "SUMMARY:Съёмка: Рекламный ролик EVERGO (смена 1 из 2)" in unfolded and "\\nОбщий сбор — 07:00" in unfolded, "")
        csv = t.js("PobubnimCallsheet.csv()")
        check(".csv: BOM, точка с запятой, строки группы",
              csv.startswith("\ufeff") and "Роль;Имя;Телефон;Вызов;Статус" in csv and
              "Оператор-постановщик;Иван Петров;+7 916 123-45-67;06:30;подтвердил" in csv, csv[:200])

        # ---------- 11. Word ----------
        b64 = t.js("(async()=>{const b=PobubnimDocx.build(document.getElementById('paper'));"
                   "const u=new Uint8Array(await b.arrayBuffer());let s='';"
                   "for(let i=0;i<u.length;i+=8192)s+=String.fromCharCode.apply(null,u.subarray(i,i+8192));"
                   "return btoa(s)})()", wait=True)
        z = zipfile.ZipFile(io.BytesIO(base64.b64decode(b64)))
        xml = z.read("word/document.xml").decode("utf-8")
        ok_xml = True
        try:
            minidom.parseString(xml.encode("utf-8"))
        except Exception as e:  # noqa: BLE001
            ok_xml = str(e)
        check(".docx распакован, document.xml — правильный XML", ok_xml is True, ok_xml)
        check(".docx: ссылка на карту кликается", 'HYPERLINK &quot;https://yandex.ru/maps/?text=' in xml, "")
        check(".docx: подпись листа — кликабельная ссылка на инструмент",
              'HYPERLINK &quot;https://pobubnim.ru/instrumenty/vyzyvnoj-list.html&quot;' in xml, "")
        check(".docx: подзаголовки цехов объединяют ячейки", '<w:gridSpan w:val="4"/>' in xml, "")
        check(".docx: расписание и группа на месте", "Сбор группы" in xml and "Иван Петров" in xml and "ВЫЗЫВНОЙ ЛИСТ" in xml, "")
        names = z.namelist()
        rels = z.read("word/_rels/document.xml.rels").decode("utf-8")
        n_img = xml.count("<w:drawing>")
        check(".docx: логотип и код маршрута вложены картинками со своими связями",
              n_img >= 1 and all(f"word/media/image{i}.png" in names and f'Target="media/image{i}.png"' in rels
                                 for i in range(1, n_img + 1)) and
              z.read("word/media/image1.png")[:8] == b"\x89PNG\r\n\x1a\n" and
              'Extension="png"' in z.read("[Content_Types].xml").decode("utf-8"), (n_img, names))
        try:
            import docx as pydocx
            doc = pydocx.Document(io.BytesIO(base64.b64decode(b64)))
            check(".docx открывается библиотекой python-docx, картинки на месте",
                  len(doc.inline_shapes) == n_img and len(doc.tables) >= 3, (len(doc.inline_shapes), len(doc.tables)))
        except ImportError:
            print("     (python-docx не установлен — проверка открытия пропущена)")

        # ---------- 12. печать ----------
        pdf = base64.b64decode(t.cmd("Page.printToPDF", printBackground=True, preferCSSPageSize=True)["result"]["data"])
        pages = len(re.findall(rb"/Type\s*/Page\b", pdf))
        check("печать: лист уходит в PDF на 1–3 страницы", 1 <= pages <= 3 and len(pdf) > 8000, f"{pages} стр., {len(pdf)} байт")

        # ---------- 13. черновик, следующая смена, очистка ----------
        time.sleep(0.9)
        t.goto(URL)
        check("черновик восстановил дату и расписание",
              t.js("document.getElementById('f-date').value") == "2026-09-14" and sched(t)[0] == ["07:00", "Сбор группы"], "")
        if cv2 is not None:
            check("логотип пережил перезагрузку", t.js("!!document.querySelector('#paper .cs-logo img')"), "")
        check("черновик восстановил группу, блоки и статусы",
              "Иван Петров" in t.paper() and "бариста" in t.paper() and t.js("PobubnimCallsheet.state().crew[0].st") == "ok", "")
        t.click("#btn-next")
        nxt = json.loads(t.js("JSON.stringify((function(){var S=PobubnimCallsheet.state();"
                              "return [document.getElementById('f-date').value,document.getElementById('f-day').value,"
                              "S.crew[0].st,S.rev,S.crew[0].who];})())"))
        check("следующая смена: дата +1, номер +1, статусы сброшены, люди те же",
              nxt == ["2026-09-15", "2 из 2", "", 0, "Иван Петров"], nxt)
        t.click(".pb-toast.on button")
        check("следующая смена отменяется", t.js("document.getElementById('f-date').value") == "2026-09-14", "")
        t.click("#btn-clear")
        time.sleep(0.6)
        check("очистка стёрла черновик и вернула подсказку",
              t.js(f"localStorage.getItem('{KEY}')") in (None, "") and t.js("!document.getElementById('starter').hidden"),
              t.js(f"localStorage.getItem('{KEY}')"))
        t.click(".pb-toast.on button")
        check("очистка отменяется кнопкой «Вернуть»", "Иван Петров" in t.paper(), "")
        t.click("#btn-clear")
        time.sleep(0.6)

        if cv2 is not None:
            check("очистка листа логотип не трогает", t.js("!!document.querySelector('#paper .cs-logo img')"), "")
            t.click("#btn-logo-del")
            check("логотип убирается своей кнопкой", t.js("!document.querySelector('#paper .cs-logo')") and
                  t.js("localStorage.getItem('pobubnim-callsheet-logo')") in (None, ""), "")
            t.click(".pb-toast.on button")
            check("и возвращается кнопкой «Вернуть»", t.js("!!document.querySelector('#paper .cs-logo img')"), "")
            t.click("#btn-logo-del")

        # пример одним кликом
        t.click("#btn-example")
        check("пример заполняет лист до полной готовности",
              t.js("document.getElementById('ready-n').textContent").startswith("Лист готов") and "Кофейня «Зерно»" in t.paper(),
              t.js("document.getElementById('ready-n').textContent + document.getElementById('ready-miss').textContent"))
        time.sleep(0.7)
        check("люди из примера в память не попадают",
              "Анна Белова" not in (t.js("localStorage.getItem('pobubnim-people-v1')") or ""), "")

        # лист первой версии: пять времён становятся расписанием
        v1 = {"locs": [{"name": "Цех", "addr": "ул. Ленина, 14", "time": "09:00"}],
              "crew": [{"role": "Гафер", "who": "Пётр", "phone": "+7 900 111-22-33", "at": "07:30"}],
              "proj": "Старый лист", "date": "2026-09-14", "day": "1", "city": "Москва", "lat": "55.7558", "lng": "37.6173",
              "tz": "3", "client": "", "call": "07:00", "go": "07:30", "start": "09:00", "lunch": "13:00", "end": "19:00",
              "notes": "Парковка во дворе."}
        t.js(f"localStorage.setItem('{KEY}', {json.dumps(json.dumps(v1, ensure_ascii=False))})")
        t.goto(URL)
        s = sched(t)
        check("лист первой версии открылся: пять времён стали строками",
              s == [["07:00", "Сбор группы"], ["07:30", "Выезд"], ["09:00", "Начало съёмки"], ["13:00", "Обед"], ["19:00", "Конец смены"]]
              and "Пётр" in t.paper() and "Парковка во дворе." in t.paper(), s)

        # ссылка на лист: черновик уезжает в адрес и возвращается
        time.sleep(0.7)
        packed = t.js("(()=>{const j=localStorage.getItem('%s');return btoa(unescape(encodeURIComponent(j)))"
                      ".replace(/\\+/g,'-').replace(/\\//g,'_').replace(/=+$/,'')})()" % KEY)
        t.js("localStorage.clear()")
        t.goto(URL + "#s=" + packed)
        check("ссылка на лист открывает его на чистом браузере", "Старый лист" in t.paper() and len(sched(t)) == 5, t.paper()[:80])
        check("десктоп без горизонтальной прокрутки",
              t.js("document.documentElement.scrollWidth <= document.documentElement.clientWidth + 1"), "")

        # ---------- 14. мобила ----------
        t.cmd("Emulation.setDeviceMetricsOverride", width=375, height=850, deviceScaleFactor=1, mobile=True)
        t.goto(URL)
        t.click("#btn-clear")
        t.click("#btn-example")
        for b in ("cast", "scenes"):
            if t.js("document.querySelector('#block-chips [data-b=\"%s\"]').getAttribute('aria-pressed')" % b) != "true":
                t.click('#block-chips [data-b="%s"]' % b)
        time.sleep(0.5)
        sw = t.js("document.documentElement.scrollWidth")
        cw = t.js("document.documentElement.clientWidth")
        check("мобила 375 без оверфлоу (и без веб-шрифтов)", sw <= cw + 1, f"{sw} > {cw}")
        check("мобильная панель на месте и знает готовность",
              t.js("getComputedStyle(document.getElementById('mbar')).display") == "flex" and
              t.js("document.getElementById('mbar-n').textContent") == "Готов", t.js("document.getElementById('mbar-n').textContent"))
        t.click("#mbar-send")
        fits = t.js("(()=>{const d=document.getElementById('send').getBoundingClientRect();return d.left>=0&&d.right<=375&&d.height<=850})()")
        check("окно рассылки помещается в телефон", t.js("document.getElementById('send').open") and fits, "")

        check("без ошибок в консоли", not t.errors, t.errors[:3])
    finally:
        t.close()

    print(("\nПРОВАЛЕНО (" + str(len(fails)) + "): " + ", ".join(fails)) if fails else "\nВсё сошлось")
    return 1 if fails else 0


if __name__ == "__main__":
    sys.exit(main())
